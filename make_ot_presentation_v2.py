import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette (TED-style, Artistic, Deep, Warm Kinfolk)
BG_CREAM = RGBColor(250, 248, 245)      # Warm Background
CARD_BG = RGBColor(255, 255, 255)       # Pure White
PRIMARY = RGBColor(44, 38, 34)          # Deep Charcoal / Warm Black
ACCENT_TERRA = RGBColor(189, 79, 44)    # Warm Terracotta
ACCENT_OLIVE = RGBColor(68, 108, 86)    # Sage / Olive Green
TEXT_MUTED = RGBColor(115, 105, 95)     # Muted Warm Gray
CARD_BORDER = RGBColor(230, 224, 215)   # Border

IMG_HEADER = "/Users/bohwanlee/.gemini/antigravity-ide/brain/9e0c0085-ca5c-46cd-ba5d-af211a1b4cc8/cooking_challenge_header_1789372226048.jpg"
IMG_VERIF = "/Users/bohwanlee/.gemini/antigravity-ide/brain/9e0c0085-ca5c-46cd-ba5d-af211a1b4cc8/cooking_verification_example_1789372243276.jpg"

blank_layout = prs.slide_layouts[6]

def set_slide_background(slide, color=BG_CREAM):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, title_text, category_text="KEYNOTE LECTURE"):
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.35))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = ACCENT_TERRA
    
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(25)
    p_title.font.bold = True
    p_title.font.color.rgb = PRIMARY

# =============================================================================
# SLIDE 1: COVER
# =============================================================================
s1 = prs.slides.add_slide(blank_layout)
set_slide_background(s1, BG_CREAM)

if os.path.exists(IMG_HEADER):
    s1.shapes.add_picture(IMG_HEADER, Inches(6.8), Inches(0.8), width=Inches(5.8), height=Inches(5.9))

tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(5.8), Inches(5.0))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "관점의 대전환 (PARADIGM SHIFT)"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_TERRA
p.space_after = Pt(16)

p2 = tf1.add_paragraph()
p2.text = "요리는 어떻게\n삶의 예술이 되는가"
p2.font.size = Pt(38)
p2.font.bold = True
p2.font.color.rgb = PRIMARY
p2.space_after = Pt(16)

p3 = tf1.add_paragraph()
p3.text = "몰입의 과학과 설명하는 힘\n[요리의 몰입과 예술 1기 온라인 OT]"
p3.font.size = Pt(15)
p3.font.color.rgb = ACCENT_OLIVE
p3.font.bold = True
p3.space_after = Pt(14)

p4 = tf1.add_paragraph()
p4.text = "1부 (40분) : 요리를 바라보는 관점 바꾸기 (인문학적 특강)\n2부 (20분) : 4주 실전 챌린지 룰 & 100% 환급 안내"
p4.font.size = Pt(13)
p4.font.color.rgb = TEXT_MUTED

# =============================================================================
# SLIDE 2: PROLOGUE
# =============================================================================
s2 = prs.slides.add_slide(blank_layout)
set_slide_background(s2, BG_CREAM)
add_header(s2, "질문 : 당신에게 '요리'는 어떤 의미입니까?", "PROLOGUE")

card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
card.fill.solid()
card.fill.fore_color.rgb = CARD_BG
card.line.color.rgb = CARD_BORDER

tb2 = s2.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.9), Inches(4.2))
tf2 = tb2.text_frame
tf2.word_wrap = True

p = tf2.paragraphs[0]
p.text = "대부분의 사람들에게 요리는 '귀찮은 가사 노동'이거나 '끼니 때우기'였습니다."
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = PRIMARY
p.space_after = Pt(24)

items2 = [
    ("일반인의 관점", "귀찮다 ➔ 배달앱을 켠다 ➔ 멍하니 스마트폰을 보며 먹는다 ➔ 더부룩함과 자괴감이 남는다."),
    ("새로운 관점", "사람이라면 누구나 매일 먹어야 하는 밥 ➔ 오롯이 내 감각에 집중하는 '몰입'의 도구로 전환한다."),
    ("오늘의 목표", "요리가 단순한 '맛'을 넘어, 내 뇌를 정화하고 삶의 질(QOL)을 극적으로 끌어올리는 경험으로 관점을 바꿉니다.")
]

for t, d in items2:
    p_it = tf2.add_paragraph()
    p_it.text = f"• {t} : "
    p_it.font.bold = True
    p_it.font.size = Pt(16)
    p_it.font.color.rgb = ACCENT_TERRA if "새로운" in t else PRIMARY
    
    run = p_it.add_run()
    run.text = d
    run.font.bold = False
    run.font.color.rgb = TEXT_MUTED
    p_it.space_after = Pt(16)

# =============================================================================
# SLIDE 3: PROBLEM - DOPAMINE OVERLOAD
# =============================================================================
s3 = prs.slides.add_slide(blank_layout)
set_slide_background(s3, BG_CREAM)
add_header(s3, "현대인의 비극 : 쉼 없는 도파민 과부하와 뇌의 피로", "PART 1. 몰입의 과학")

problems = [
    ("01. 스마트폰 중독", "아침 눈뜰 때부터 잠들 때까지\n숏폼, 릴스, 알림에 뇌가 절여져\n단 10분도 온전히 집중하지 못함"),
    ("02. 감각의 상실", "음식을 먹으면서도 유튜브를 보느라\n내가 지금 무슨 맛을 느끼는지,\n배가 부른지조차 인지하지 못함"),
    ("03. 만성 무기력", "자극적인 도파민이 쏟아진 뒤\n찾아오는 공허함과 무력감.\n내 삶을 내가 통제하지 못한다는 결핍")
]

for i, (tit, desc) in enumerate(problems):
    card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.95), Inches(1.8), Inches(3.7), Inches(5.0))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    
    tb = s3.shapes.add_textbox(Inches(1.1 + i * 3.95), Inches(2.2), Inches(3.1), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = tit
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(20)
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(24)
    
    p3 = tf.add_paragraph()
    p3.text = "➔ '진짜 몰입'이 절실한 이유"
    p3.font.size = Pt(13)
    p3.font.bold = True
    p3.font.color.rgb = ACCENT_TERRA

# =============================================================================
# SLIDE 4: CHEF'S INTERVIEW - BEYOND TASTE
# =============================================================================
s4 = prs.slides.add_slide(blank_layout)
set_slide_background(s4, BG_CREAM)
add_header(s4, "세계적인 셰프들은 왜 '맛 이상의 어떤 것'을 말할까요?", "PART 1. 몰입의 과학")

card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
card.fill.solid()
card.fill.fore_color.rgb = CARD_BG
card.line.color.rgb = CARD_BORDER

tb4 = s4.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.9), Inches(4.2))
tf4 = tb4.text_frame
tf4.word_wrap = True

p = tf4.paragraphs[0]
p.text = "“요리는 음식을 만드는 행위가 아니다. 나를 잊고 재료와 하나가 되는 명상이다.”"
p.font.size = Pt(21)
p.font.bold = True
p.font.color.rgb = ACCENT_TERRA
p.space_after = Pt(20)

items4 = [
    ("미슐랭 셰프들의 공통된 인터뷰", "그들은 하나같이 말합니다. '진짜 요리는 맛을 넘어선 무아지경의 몰입(Flow)이다.' 칼질의 규칙적인 리듬, 팬 위에서 변해가는 식재료의 색과 향에 온 신경이 쏠릴 때 시간은 멈춥니다."),
    ("움직이는 선(Zen)과 명상", "가만히 앉아 명상하는 것은 어렵지만, 요리는 오감(시각, 청각, 후각, 촉각, 미각)이 강제로 동원되기 때문에 스마트폰과 잡념이 즉각적으로 차단됩니다."),
    ("우리가 추구할 요리", "우리는 주방 노동자가 아닙니다. 내 삶의 주권을 되찾고, 하루 중 유일하게 뇌가 맑아지는 '몰입의 시간'을 누리는 것입니다.")
]

for t, d in items4:
    p_it = tf4.add_paragraph()
    p_it.text = f"• {t} : "
    p_it.font.bold = True
    p_it.font.size = Pt(15)
    p_it.font.color.rgb = PRIMARY
    
    run = p_it.add_run()
    run.text = d
    run.font.bold = False
    run.font.color.rgb = TEXT_MUTED
    p_it.space_after = Pt(15)

# =============================================================================
# SLIDE 5: FLOW TO HABIT (몰입에서 습관으로)
# =============================================================================
s5 = prs.slides.add_slide(blank_layout)
set_slide_background(s5, BG_CREAM)
add_header(s5, "요리의 몰입이 '습관'으로 넘어갈 때 일어나는 기적", "PART 1. 몰입의 과학")

flow_steps = [
    ("1단계. 일시적 몰입", "주방에서 30분간\n스마트폰을 끄고\n온전한 도파민 디톡스 경험"),
    ("2단계. 루틴의 형성", "주 3회 규칙적인 요리\n재료를 다루는 손길이\n익숙해지고 불안이 감소"),
    ("3단계. 삶의 질 수직 상승", "내 손으로 나를 대접하는 자존감\n맑은 정신, 건강한 신체,\n삶의 통제권 완벽 회복!")
]

for i, (st, desc) in enumerate(flow_steps):
    card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.95), Inches(2.0), Inches(3.7), Inches(4.5))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    
    tb = s5.shapes.add_textbox(Inches(1.1 + i * 3.95), Inches(2.3), Inches(3.1), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = f"PHASE 0{i+1}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TERRA
    p.space_after = Pt(10)
    
    p2 = tf.add_paragraph()
    p2.text = st
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = PRIMARY
    p2.space_after = Pt(16)
    
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED

# =============================================================================
# SLIDE 6: WHAT IS ART? (예술이란 무엇인가)
# =============================================================================
s6 = prs.slides.add_slide(blank_layout)
set_slide_background(s6, BG_CREAM)
add_header(s6, "그렇다면 '예술'이란 진정 무엇일까요?", "PART 2. 설명하는 힘")

card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
card.fill.solid()
card.fill.fore_color.rgb = CARD_BG
card.line.color.rgb = CARD_BORDER

tb6 = s6.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.9), Inches(4.2))
tf6 = tb6.text_frame
tf6.word_wrap = True

p = tf6.paragraphs[0]
p.text = "벽에 테이프로 붙인 바나나 하나가 1억 5천만 원에 팔리는 이유"
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = PRIMARY
p.space_after = Pt(20)

art_items = [
    ("형태가 아니라 '설명'이다", "마우리치오 카텔란의 바나나 작품 <코미디언>, 뒤샹의 변기 <샘>. 그들이 대단한 솜씨로 빚어낸 형태입니까? 아닙니다. 거기에 붙은 작가의 철학과 '설명'이 평범한 사물을 예술로 승격시킨 것입니다."),
    ("예술을 잘한다는 것의 정의", "예술을 잘하는 사람은 그림을 잘 그리는 사람이 아닙니다. '자신이 만든 결과물에 왜 이런 의도를 담았는지 언어로 설득하고 납득시키는 사람'입니다."),
    ("요리와 예술의 일치", "설명 없는 요리는 그저 '배고파서 때워 넣는 칼로리'에 불과합니다. 하지만 나의 의도와 이유가 설명되는 순간, 그 요리는 나만의 식탁에 오른 '예술 작품'이 됩니다.")
]

for t, d in art_items:
    p_it = tf6.add_paragraph()
    p_it.text = f"• {t} : "
    p_it.font.bold = True
    p_it.font.size = Pt(15)
    p_it.font.color.rgb = ACCENT_TERRA if "정의" in t else PRIMARY
    
    run = p_it.add_run()
    run.text = d
    run.font.bold = False
    run.font.color.rgb = TEXT_MUTED
    p_it.space_after = Pt(15)

# =============================================================================
# SLIDE 7: POWER OF EXPLANATION IN REAL LIFE (일상의 무기)
# =============================================================================
s7 = prs.slides.add_slide(blank_layout)
set_slide_background(s7, BG_CREAM)
add_header(s7, "일상과 일에서도 '설명하는 힘'이 내 가치를 결정합니다", "PART 2. 설명하는 힘")

# 2 Columns (Before vs After / Comparison)
col1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
col1.fill.solid()
col1.fill.fore_color.rgb = CARD_BG
col1.line.color.rgb = CARD_BORDER

tb_c1 = s7.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.3))
tf_c1 = tb_c1.text_frame
tf_c1.word_wrap = True

p = tf_c1.paragraphs[0]
p.text = "❌ 설명하지 못하는 사람"
p.font.size = Pt(19)
p.font.bold = True
p.font.color.rgb = RGBColor(180, 60, 60)
p.space_after = Pt(16)

c1_text = [
    ("직장 보고서", "열심히 밤새서 일했지만 '그냥 데이터 정리했습니다'라고 말해 노력의 가치를 폄하당함."),
    ("인간관계", "좋은 의도로 한 행동이지만 맥락을 설명하지 못해 오해를 사거나 당연하게 여겨짐."),
    ("나의 요리", "'배고파서 대충 양파랑 계란 볶았어.' ➔ 영혼 없는 식사, 자존감 정체.")
]
for t, d in c1_text:
    p_i = tf_c1.add_paragraph()
    p_i.text = f"• {t}\n  {d}\n"
    p_i.font.size = Pt(13)
    p_i.font.color.rgb = TEXT_MUTED

# Col 2
col2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
col2.fill.solid()
col2.fill.fore_color.rgb = CARD_BG
col2.line.color.rgb = ACCENT_TERRA

tb_c2 = s7.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.1), Inches(4.3))
tf_c2 = tb_c2.text_frame
tf_c2.word_wrap = True

p = tf_c2.paragraphs[0]
p.text = "✔ 설명하는 힘을 기른 사람"
p.font.size = Pt(19)
p.font.bold = True
p.font.color.rgb = ACCENT_TERRA
p.space_after = Pt(16)

c2_text = [
    ("직장 보고서", "'이 데이터의 변화가 우리 사업에 어떤 기회인지' 의도와 스토리를 입혀 탁월한 평가를 받음."),
    ("인간관계", "나의 생각과 마음의 이유를 상대가 납득할 수 있게 부드럽게 전달하여 신뢰를 얻음."),
    ("나의 요리", "'오늘 힘든 나를 위해 1시간 동안 정성껏 카라멜라이징한 양파 수프야.' ➔ 나를 극진히 대접하는 예술.")
]
for t, d in c2_text:
    p_i = tf_c2.add_paragraph()
    p_i.text = f"• {t}\n  {d}\n"
    p_i.font.size = Pt(13)
    p_i.font.color.rgb = PRIMARY

# =============================================================================
# SLIDE 8: THE BRIDGE (선언)
# =============================================================================
s8 = prs.slides.add_slide(blank_layout)
set_slide_background(s8, PRIMARY)

tb8 = s8.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(4.5))
tf8 = tb8.text_frame
tf8.word_wrap = True

p = tf8.paragraphs[0]
p.text = "THE TURNING POINT"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_TERRA
p.space_after = Pt(20)

p2 = tf8.add_paragraph()
p2.text = "“그렇기 때문에,\n우리는 이 4주 챌린지를 시작합니다.”"
p2.font.size = Pt(38)
p2.font.bold = True
p2.font.color.rgb = RGBColor(255, 255, 255)
p2.space_after = Pt(24)

p3 = tf8.add_paragraph()
p3.text = "단순한 요리 레시피 모임이 아닙니다.\n1) 주방에서 스마트폰을 끄고 누리는 '절대적 몰입 (도파민 디톡스)'\n2) 내 요리와 행동에 스토리를 입히는 '설명하는 힘'\n\n이 두 가지 무기를 손에 쥐고 삶의 질을 완전히 다른 차원으로 끌어올리는 28일입니다."
p3.font.size = Pt(16)
p3.font.color.rgb = RGBColor(215, 205, 195)

# =============================================================================
# SLIDE 9: CURRICULUM OVERVIEW (2부 시작)
# =============================================================================
s9 = prs.slides.add_slide(blank_layout)
set_slide_background(s9, BG_CREAM)
add_header(s9, "4주 커리큘럼 : 1주 1재료 3요리 원칙", "PART 3. 실전 가이드 (20분)")

curriculums = [
    ("1주차", "양파 (Onion)", "단단한 껍질을 벗겨내고\n불을 만나 달콤해지는 시간", "카라멜라이징, 양파스프,\n양파피클, 양파덮밥 등"),
    ("2주차", "계란 (Egg)", "가장 단순하지만\n무궁무진한 변주의 미학", "달걀말이, 수란 토스트,\n스페니시 오믈렛, 달걀장 등"),
    ("3주차", "토마토 (Tomato)", "붉은 생명력과 산미,\n깊은 풍미를 끌어내는 몰입", "토마토 마리네이드, 파스타,\n토마토 달걀볶음, 토마토스프 등"),
    ("4주차", "감자 (Potato)", "투박한 흙에서 건져 올린\n든든함과 따뜻한 위로", "감자 뇨끼, 감자채전,\n메쉬드 포테이토, 감자조림 등")
]

for i, (wk, ing, theme, ex) in enumerate(curriculums):
    card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 2.95), Inches(1.8), Inches(2.8), Inches(4.8))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    
    tb = s9.shapes.add_textbox(Inches(1.0 + i * 2.95), Inches(2.0), Inches(2.4), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = wk
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TERRA
    p.space_after = Pt(8)
    
    p2 = tf.add_paragraph()
    p2.text = ing
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = PRIMARY
    p2.space_after = Pt(14)
    
    p3 = tf.add_paragraph()
    p3.text = theme
    p3.font.size = Pt(13)
    p3.font.color.rgb = ACCENT_OLIVE
    p3.space_after = Pt(16)
    
    p4 = tf.add_paragraph()
    p4.text = f"[추천 요리 예시]\n{ex}"
    p4.font.size = Pt(12)
    p4.font.color.rgb = TEXT_MUTED

# =============================================================================
# SLIDE 10: RECIPE NOTEBOOK DOCTRINE
# =============================================================================
s10 = prs.slides.add_slide(blank_layout)
set_slide_background(s10, BG_CREAM)
add_header(s10, "왜 스마트폰을 끄고 '종이 레시피북'에 써야 하는가?", "PART 3. 실전 가이드")

tb10 = s10.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(6.5), Inches(5.0))
tf10 = tb10.text_frame
tf10.word_wrap = True

p = tf10.paragraphs[0]
p.text = "스마트폰이 주방에 들어오는 순간, 몰입은 깨집니다."
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = ACCENT_TERRA
p.space_after = Pt(16)

rules10 = [
    ("스마트폰 프리존 (Free Zone)", "유튜브나 릴스를 보며 요리하면 음식에 집중할 수 없습니다. 스마트폰은 거실에 두고 주방에는 오직 종이 노트만 들고 들어갑니다."),
    ("손글씨로 뇌에 새기는 각인", "직접 펜으로 적어본 재료와 조리 순서는 뇌에 깊이 각인되어 평생 내 기술이 됩니다."),
    ("나만의 아틀리에 보물", "4주 뒤, 손때 묻고 기름방울이 살짝 튄 세상에 단 하나뿐인 [12가지 나만의 레시피북]이 완성됩니다.")
]

for title, desc in rules10:
    p_item = tf10.add_paragraph()
    p_item.text = f"✔ {title}\n"
    p_item.font.bold = True
    p_item.font.size = Pt(15)
    p_item.font.color.rgb = PRIMARY
    
    run = p_item.add_run()
    run.text = desc
    run.font.bold = False
    run.font.size = Pt(13)
    run.font.color.rgb = TEXT_MUTED
    p_item.space_after = Pt(14)

if os.path.exists(IMG_VERIF):
    s10.shapes.add_picture(IMG_VERIF, Inches(7.5), Inches(1.8), width=Inches(5.0), height=Inches(4.8))

# =============================================================================
# SLIDE 11: VERIFICATION RULES
# =============================================================================
s11 = prs.slides.add_slide(blank_layout)
set_slide_background(s11, BG_CREAM)
add_header(s11, "카톡방 주간 인증 3단계 (사진 + 설명)", "PART 3. 실전 가이드")

steps = [
    ("STEP 1", "주 3회 요리 & 레시피북 작성", "매주 지정 식재료로 3번 요리하고,\n종이 노트에 레시피와 의도를 손글씨로 기록합니다."),
    ("STEP 2", "레시피북 사진 3장 촬영", "직접 쓴 레시피북 페이지를 촬영합니다.\n(완성된 요리 사진을 함께 올리면 환상적!)"),
    ("STEP 3", "카톡방에 '요리의 설명' 업로드", "이 요리를 만든 이유, 생각, 추억을\n언어로 풀어내어 단톡방 크루들과 공유합니다.")
]

for i, (st, title, desc) in enumerate(steps):
    card = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.95), Inches(1.8), Inches(3.7), Inches(2.6))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    
    tb = s11.shapes.add_textbox(Inches(1.0 + i * 3.95), Inches(2.0), Inches(3.3), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = st
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TERRA
    p.space_after = Pt(6)
    
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = PRIMARY
    p2.space_after = Pt(10)
    
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(13)
    p3.font.color.rgb = TEXT_MUTED

# Example Box
box_ex = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.7), Inches(11.7), Inches(2.1))
box_ex.fill.solid()
box_ex.fill.fore_color.rgb = RGBColor(243, 239, 233)
box_ex.line.color.rgb = CARD_BORDER

tb_ex = s11.shapes.add_textbox(Inches(1.1), Inches(4.9), Inches(11.1), Inches(1.7))
tf_ex = tb_ex.text_frame
tf_ex.word_wrap = True

pe1 = tf_ex.paragraphs[0]
pe1.text = "💬 [카톡 인증 양식 예시]"
pe1.font.bold = True
pe1.font.size = Pt(14)
pe1.font.color.rgb = ACCENT_TERRA
pe1.space_after = Pt(6)

pe2 = tf_ex.add_paragraph()
pe2.text = "📷 사진 : 이번 주 작성한 레시피북 3페이지 촬영 사진\n✍️ 요리의 설명 : \n1. [양파 카라멜라이징 수프] 껍질을 까며 복잡한 생각을 비워내고, 40분 동안 불 앞을 지키며 인내의 단맛을 맛본 요리\n2. [양파 달걀 덮밥] 배달 대신 10분 만에 따뜻하게 나를 위로해 준 든든한 한 끼"
pe2.font.size = Pt(12)
pe2.font.color.rgb = PRIMARY

# =============================================================================
# SLIDE 12: REFUND SYSTEM
# =============================================================================
s12 = prs.slides.add_slide(blank_layout)
set_slide_background(s12, BG_CREAM)
add_header(s12, "100% 환급 시스템 : 배달비 아끼고 전액 다시 찾아가기!", "PART 3. 실전 가이드")

refund_cards = [
    ("100% 완주 달성", "12회 완벽 인증", "50,000원", "전액 100% 환급!", ACCENT_TERRA),
    ("90% 이상 달성", "11회 인증", "35,000원", "70% 환급", ACCENT_OLIVE),
    ("75% 이상 달성", "10회 인증", "25,000원", "50% 환급", PRIMARY)
]

for i, (title, sub, money, rate, color) in enumerate(refund_cards):
    card = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.95), Inches(1.8), Inches(3.7), Inches(3.5))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    
    tb = s12.shapes.add_textbox(Inches(1.0 + i * 3.95), Inches(2.2), Inches(3.3), Inches(2.8))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(8)
    
    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(20)
    
    p3 = tf.add_paragraph()
    p3.text = money
    p3.font.size = Pt(28)
    p3.font.bold = True
    p3.font.color.rgb = color
    p3.space_after = Pt(6)
    
    p4 = tf.add_paragraph()
    p4.text = rate
    p4.font.size = Pt(15)
    p4.font.bold = True
    p4.font.color.rgb = color

tb_note = s12.shapes.add_textbox(Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.2))
tf_note = tb_note.text_frame
tf_note.word_wrap = True
pn = tf_note.paragraphs[0]
pn.text = "💡 5만 원은 비용이 아니라 내 완주 의지를 지키는 '보증금'입니다.\n20명 전원이 100% 환급받아 가실 수 있도록 단톡방에서 끝까지 함께 응원하겠습니다!"
pn.font.size = Pt(14)
pn.font.bold = True
pn.font.color.rgb = ACCENT_TERRA

# =============================================================================
# SLIDE 13: PREPARATION WEEK MISSIONS
# =============================================================================
s13 = prs.slides.add_slide(blank_layout)
set_slide_background(s13, BG_CREAM)
add_header(s13, "준비 주간 미션 (9/17 ~ 9/20) : 나만의 도구를 갖추라!", "PART 3. 실전 가이드")

card = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
card.fill.solid()
card.fill.fore_color.rgb = CARD_BG
card.line.color.rgb = CARD_BORDER

tb13 = s13.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.9), Inches(4.2))
tf13 = tb13.text_frame
tf13.word_wrap = True

p = tf13.paragraphs[0]
p.text = "본 챌린지 시작(9/21 월) 전까지 아래 3가지를 완료해 주세요!"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = ACCENT_TERRA
p.space_after = Pt(20)

m_items = [
    ("미션 1. 마음에 쏙 드는 '종이 레시피북(노트)' 마련하기", "문구점이나 다이소에서 마음에 드는 무선/유선 노트를 하나 구입하세요. 표지에 [요리의 몰입과 예술 1기 - 나의 이름]을 손글씨로 적어봅니다."),
    ("미션 2. 1주차 식재료 '양파' 장보기", "9/21(월) 첫 요리를 시작하기 전, 신선한 양파 1망을 사서 냉장고나 서늘한 곳에 준비해 둡니다."),
    ("미션 3. 양파로 해보고 싶은 3가지 요리 스케치하기", "노트 첫 페이지에 이번 주에 도전해보고 싶은 3가지 양파 요리 이름을 미리 손으로 적어봅니다.")
]

for title, desc in m_items:
    p_item = tf13.add_paragraph()
    p_item.text = f"✔ {title}\n"
    p_item.font.bold = True
    p_item.font.size = Pt(16)
    p_item.font.color.rgb = PRIMARY
    
    run = p_item.add_run()
    run.text = desc
    run.font.bold = False
    run.font.size = Pt(14)
    run.font.color.rgb = TEXT_MUTED
    p_item.space_after = Pt(14)

# =============================================================================
# SLIDE 14: CLOSING & Q&A
# =============================================================================
s14 = prs.slides.add_slide(blank_layout)
set_slide_background(s14, BG_CREAM)
add_header(s14, "주방은 나를 대접하는 가장 작은 아틀리에입니다", "FINISH")

tb14 = s14.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.5))
tf14 = tb14.text_frame
tf14.word_wrap = True

p = tf14.paragraphs[0]
p.text = "Q & A 및 설레는 1기 시작 단체 캡처 타임"
p.font.size = Pt(30)
p.font.bold = True
p.font.color.rgb = ACCENT_TERRA
p.space_after = Pt(20)

p2 = tf14.add_paragraph()
p2.text = "• 궁금하신 점이나 룰 관련 질문을 편하게 남겨주세요.\n" \
          "• 카메라를 켜고, 4주간 함께할 크루들과 첫 단체 기념 캡처를 남깁니다!\n\n" \
          "“우리는 4주 뒤, 12가지 나만의 레시피와 설명하는 힘을 품은 완전히 새로운 나를 만납니다.”"
p2.font.size = Pt(17)
p2.font.color.rgb = PRIMARY

output_path = "/Users/bohwanlee/Desktop/antigravity 폴더/요리의_몰입과_예술_1기_OT.pptx"
prs.save(output_path)
print(f"Upgraded PPTX Successfully created at: {output_path}")
